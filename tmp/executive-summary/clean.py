from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
src=Document('output/documents/Curbside_Executive_Summary_Template.docx')
sections=[]
for t in src.tables:
 for row in t.rows:
  if len(row.cells)!=2:continue
  texts=[p.text for p in row.cells[0].paragraphs if p.text]
  if texts:sections.append((texts[0],texts[1:]))
D=Document();s=D.sections[0];s.top_margin=s.bottom_margin=Inches(.7);s.left_margin=s.right_margin=Inches(.8)
s.page_width=Inches(8.5);s.page_height=Inches(11)
style=D.styles['Normal'];style.font.name='Calibri';style.font.size=Pt(12);style.paragraph_format.space_after=Pt(9);style.paragraph_format.line_spacing=1.1
for name,size in [('Title',27),('Heading 1',16)]:
 st=D.styles[name];st.font.name='Calibri';st.font.size=Pt(size);st.font.color.rgb=RGBColor(0,0,0);st.paragraph_format.space_before=Pt(13);st.paragraph_format.space_after=Pt(7)
for e in D.styles.element.xpath('.//w:pBdr'):e.getparent().remove(e)
D.add_paragraph('Curbside',style='Title')
p=D.add_paragraph('Non-Confidential Executive Summary');p.paragraph_format.space_after=Pt(5)
p=D.add_paragraph('curbsidescore.com  |  [Contact name and cell phone]');p.paragraph_format.space_after=Pt(15)
for r in p.runs:r.font.size=Pt(10);r.font.color.rgb=RGBColor.from_string('555555')
for heading,pars in sections:
 if heading.startswith('Competitors'):heading='Competitors'
 p=D.add_paragraph(heading,style='Heading 1')
 if heading in ['Product/Services','Customers','Competitors','Management Team']:p.paragraph_format.page_break_before=True
 for text in pars:
  if heading=='Management Team' and text.startswith('[Founder'):
   text='[Founder name, degree program, relevant experience, and accomplishments.] Academic advisors: [names or none]. External advisors: [names or none]. The team has developed the initial evaluation product and will build on that foundation through customer research and partner pilots.'
  if heading=='Business Model' and text.startswith('Current revenue'):
   text+=' Capital received and sources: [confirm amount and sources, or none].'
  if heading=='Goals' and text.startswith('At the Hurricane'):
   text=text.replace('early funding','[confirm funding amount] in early funding')
  D.add_paragraph(text)
D.add_paragraph('Financials ($ US)',style='Heading 1')
p=D.add_paragraph('Illustrative consumer projections assume 500, 2,000, and 5,000 annual payers at $12 average revenue per payer. Costs exclude founder salary, taxes, and additional institutional development. Partnership and licensing revenue will be modeled after pilots establish terms. Historical figures require confirmation.')
for r in p.runs:r.font.size=Pt(10)
t=D.add_table(rows=0,cols=6);t.style='Light Shading Accent 1'
for row in src.tables[-1].rows:
 cells=t.add_row().cells
 for c,old in zip(cells,row.cells):
  c.text=old.text
  for p in c.paragraphs:
   p.paragraph_format.space_after=Pt(5)
   for r in p.runs:r.font.size=Pt(9)
for p in src.paragraphs:
 if p.text.startswith('[1]'):
  x=D.add_paragraph(p.text);x.paragraph_format.space_before=Pt(10)
  for r in x.runs:r.font.size=Pt(8);r.font.color.rgb=RGBColor.from_string('555555')
f=s.footer.paragraphs[0];f.alignment=2
r=f.add_run('Curbside  |  ');r.font.size=Pt(9)
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');f._p.append(fld)
D.save('output/documents/Curbside_Executive_Summary.docx')
