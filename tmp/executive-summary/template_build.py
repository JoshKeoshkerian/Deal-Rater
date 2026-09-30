from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
D=Document(); sec=D.sections[0]
sec.top_margin=sec.bottom_margin=Inches(.45);sec.left_margin=sec.right_margin=Inches(.4)
sec.page_width=Inches(8.5);sec.page_height=Inches(11)
for name in ['Normal','Title','Heading 1','Heading 2']:
 st=D.styles[name];st.font.name='Calibri';st.font.color.rgb=RGBColor(0,0,0)
 st.font.size=Pt(10.5);st.paragraph_format.space_after=Pt(6);st.paragraph_format.line_spacing=1.03
for e in D.styles.element.xpath('.//w:pBdr'):e.getparent().remove(e)
def para(cell,text,bold=False):
 p=cell.add_paragraph();p.paragraph_format.space_after=Pt(7)
 r=p.add_run(text);r.bold=bold
 return p
def header(first=False):
 if not first:D.add_page_break()
 t=D.add_table(rows=1,cols=2);t.autofit=False;t.columns[0].width=Inches(3.4);t.columns[1].width=Inches(4.3)
 t.cell(0,0).paragraphs[0].add_run().add_picture('extension/icons/icon128.png',width=Inches(.46))
 p=t.cell(0,1).paragraphs[0];p.alignment=2;r=p.add_run('Non-Confidential Executive Summary');r.bold=True;r.font.size=Pt(14)
 p=D.add_paragraph('Curbside',style='Title');p.runs[0].bold=True;p.runs[0].font.size=Pt(13)
 pr=p._p.get_or_add_pPr();b=OxmlElement('w:pBdr');bot=OxmlElement('w:bottom');bot.set(qn('w:val'),'single');bot.set(qn('w:sz'),'8');bot.set(qn('w:color'),'AAAAAA');b.append(bot);pr.append(b)
 t=D.add_table(rows=0,cols=2);t.autofit=False;t.columns[0].width=Inches(5.45);t.columns[1].width=Inches(2.25)
 return t
def row(t,heading,paragraphs,fields=[]):
 cells=t.add_row().cells
 for c in cells:
  c.paragraphs[0].paragraph_format.space_after=Pt(0);c.paragraphs[0].paragraph_format.space_before=Pt(0);c.paragraphs[0].paragraph_format.line_spacing=Pt(1);c.paragraphs[0].add_run().font.size=Pt(1)
 para(cells[0],heading,True)
 for v in paragraphs:para(cells[0],v)
 for label,value in fields:
  para(cells[1],label,True);para(cells[1],value)
 return cells
T=header(True)
row(T,'Business Summary',[
'Curbside is building a used-car buying platform that guides people through the entire private-party purchase, from comparing vehicles and understanding risks to negotiating, arranging inspections, and connecting with financing and secure transaction services. Our mission is to make this process clearer and more accessible for people who depend on reliable transportation.',
'Our first product is a Chrome extension that analyzes Facebook Marketplace car listings and produces a score out of 100, supported by comparable asking prices, vehicle risk indicators, and negotiation guidance. It introduces Curbside at an early decision point: deciding whether a vehicle deserves further attention.',
'Our strategy is to grow from this evaluation product into a connected buying experience. We plan to add service-provider partnerships and offer our technology through software subscriptions and licensing to credit unions and automotive platforms. Our objectives are to validate consumer value, connect useful purchase services, and establish a path to recurring business revenue.'
],[('Contact Person & Cell Phone:','[Name and cell phone]')])
row(T,'Customer Problem',[
'Private-party used-car buyers face a major financial decision with incomplete information and limited support. Listings can omit important details, asking prices are difficult to compare, and buyers often lack the experience to recognize warning signs or negotiate effectively. A seemingly affordable vehicle can become an expensive mistake.',
'The problem continues after a promising vehicle is found. Buyers must piece together pricing research, vehicle history, inspections, financing, and purchase logistics across separate services. Students, first-time buyers, and budget-conscious households may have limited time, money, or experience to manage that process.',
'Curbside aims to turn scattered information into clear, actionable guidance throughout the purchase. Consumers would pay for faster screening, understandable risk explanations, and help deciding what to do next. Free evaluations allow them to experience the value before purchasing additional checks or a shopping-period subscription.',
'Willingness to pay will be tested through completed purchases, interviews, and repeat use. For future business customers, the proposed value is providing useful buying support within an existing member or customer relationship. Paid pilots will test that proposition separately from consumer demand.'
],[('Website:','curbsidescore.com')])
T=header()
row(T,'Product/Services',[
'The current product evaluates the listing a buyer is viewing, compares relevant asking prices, and explains vehicle risk indicators, listing concerns, and negotiation context. Saved evaluations help buyers revisit a shortlist. The score is supported by explanations so buyers can understand the factors behind it. Asking-price comparisons are not verified transaction values, and listing analysis does not replace an independent inspection.',
'Our planned platform would carry this guidance through the purchase. A buyer could compare vehicles, prepare questions for the seller, arrange an inspection through a provider, and connect with financing and secure transaction services. These capabilities are future development objectives; provider relationships have not been confirmed.',
'We also intend to offer evaluation and buying tools to credit unions and automotive platforms through subscriptions, licensing, and integrations. A credit union could provide members with vehicle research support alongside its financing process. An automotive platform could incorporate Curbside guidance into its existing customer experience.',
'The implemented foundation includes capture, evaluation, accounts, saved results, and payment integration code. An internal review reported 15.3% median absolute percentage error across 4,133 held-out asking-price predictions. Overall verdict agreement was 38% across 72 labels from one reviewer. These are preliminary internal results; broader independent validation remains necessary. [1]',
'Our intellectual property position centers on original software, evaluation methods, and proprietary know-how. [Confirm copyright ownership, trademark or patent status, and third-party licensing.] Future defensibility would also depend on permitted outcome feedback and partner integrations.',
'Operating priorities include platform permissions, privacy, data rights, and clear responsibility for partner services. Future financing and transaction arrangements require specialist review before implementation. Referral compensation should be disclosed and separated from evaluation scores. No measured environmental benefit is claimed.'
],[('University:','[University]')])
row(T,'Target Market',[
'Curbside initially targets U.S. private-party buyers using Facebook Marketplace on desktop Chrome. Cox Automotive’s December 2025 forecast projected 38.3 million U.S. used-vehicle sales in 2026, down 0.9% from 2025. This is a broad industry benchmark, not a count of eligible users. [2]',
'An illustrative reachable segment of 5% of that volume would equal 1.915 million shopping episodes; the percentage requires validation. The broader opportunity includes credit unions and automotive platforms buying software, plus services used during a purchase. Institutional contract values and partner revenue potential will be sized through pilot discussions, rather than treated as established demand.'
],[('Academic Advisors:','[Names or none]')])
T=header()
row(T,'Customers',[
'Our priority consumers are students, first-time buyers, and budget-conscious households actively searching for reliable transportation. They need understandable comparisons, practical questions, and help coordinating the steps between finding a vehicle and completing a purchase.',
'Future business customers are credit unions seeking member-facing buying tools and automotive platforms seeking integrated evaluation and purchase guidance. Inspection, financing, and transaction providers are potential service partners. These groups have different needs and will be validated through separate interviews and pilots.'
])
row(T,'Sales/Marketing Strategy',[
'Initial distribution will use the Chrome Web Store and Curbside website. Product demonstrations, educational content, university networks, appropriate automotive communities, and referrals will show how an evaluation supports a real purchase decision. Paid acquisition will expand only after conversion and customer acquisition costs justify it.',
'An illustrative first-year funnel targets 50,000 qualified visits, 10,000 activated shoppers, and 500 paying customers. This assumes 20% activation and 5% paid conversion. Against the illustrative 1.915 million reachable shopping episodes, those targets equal approximately 0.52% activation penetration and 0.026% paid penetration, assuming one episode per customer.',
'In parallel, we plan to approach inspection providers, credit unions, and automotive platforms for focused pilots. A successful pilot should establish customer use, delivery cost, and a clear purchasing decision. Business distribution would extend reach beyond shoppers installing the extension directly; no signed distribution agreements are assumed.'
],[('External Advisors (if any):','[Names or none]')])
row(T,'Business Model',[
'The first revenue stream is consumer freemium access. Current pricing provides ten free checks, ten additional checks for $7.99, twenty for $11.99, or unlimited evaluations for $16.99 per month. Paid packs do not expire. This accommodates a concentrated car-shopping period without assuming a year-long consumer subscription. [6]',
'The second planned stream is partnership revenue from inspection, financing, and secure transaction services, subject to appropriate agreements and review. The third is recurring software subscriptions and licensing to credit unions and automotive platforms. Business pricing will reflect usage, integration, support, and demonstrated value.',
'Current revenue and monthly burn: [confirm actual amounts]. The illustrative consumer operating budget is $12,000 in 2026, $24,000 in 2027, and $42,000 in 2028, averaging $1,000, $2,000, and $3,500 monthly. With the revenue assumptions below, annual net cash results are negative $6,000, break-even, and positive $18,000 before founder compensation and taxes. Institutional development requires a separate pilot budget.'
],[('Capital Received and Source(s):','[Amount and sources, or none]')])
T=header()
row(T,'Competitors (do not state that you have no competition):',[
'Kelley Blue Book provides vehicle valuations, including private-party values. CarGurus offers valuation tools and deal ratings. CARFAX provides vehicle-history reports. These established brands address important parts of a buyer’s research and bring recognition and data resources that Curbside must respect. [3–5]',
'Manual listing comparisons, advice from knowledgeable friends, and independent mechanics are additional substitutes or complements. As Curbside expands, individual inspection, financing, and transaction services will also compete for portions of the buying experience, even where some providers could become partners.',
'Free research tools create a meaningful price hurdle. Curbside’s paid consumer offer, currently $7.99 to $16.99, must demonstrate useful listing-specific guidance and convenience. Our proposed market position is a connected buying experience: helping a customer interpret a listing and then take practical steps toward purchase.',
'Our intended strength is continuity across decisions that buyers currently coordinate themselves. The current extension provides an accessible entry point; planned services would extend support through inspection, financing, and completion. We do not claim superior valuation accuracy or vehicle-history coverage without supporting evidence.',
'Competitor shares of this specific private-party buying-support market are not established in the available evidence. We therefore do not assign unsupported percentages. Pilot research will compare usability, customer trust, and willingness to pay against the tools buyers already use.'
],[('Capital Seeking *:','[Confirm amount]'),('Use(s) of funding:','Evaluation reliability; buyer research and acquisition; a focused service or institutional pilot; specialist review and contingency.')])
row(T,'Competitive Advantage',[
'Curbside combines evaluation technology with a buyer-centered process that begins where a shopper finds a car. Its broader vision connects explanations with the next useful action, so the buyer does not have to rebuild context at each stage of the purchase.',
'The intended long-term advantage has three sources: better guidance informed by permitted buyer feedback, integrations that connect relevant services, and distribution through trusted organizations. For example, an inspection outcome could inform future evaluation improvements, while a credit union could introduce the tools to members already considering a vehicle purchase.',
'These advantages remain to be developed. The current feature set is reproducible, and no exclusive network or defensible data advantage is claimed. Sustainable differentiation will depend on reliable results, useful integrations, strong relationships, and customer trust. Partner compensation must not determine the evaluation score.'
],[('* Please do not include a suggested company valuation','')])
T=header()
row(T,'Management Team',[
'[Founder name, degree program, relevant experience, education, past performance, start-up history, and prizes.] The implemented product provides a foundation for execution; the final submission should connect that work to the team members responsible for it.',
'The venture needs capabilities in software development, automotive evaluation, buyer research, partnerships, and business sales. Early gaps can be addressed with advisors and targeted specialist engagements. Institutional deployment will add security, integration, and support requirements. [Identify current strengths and priority hires.]',
'We will use customer interviews, independent evaluation reviews, and pilot feedback to guide product and commercial decisions. This keeps the team responsive to evidence as the business grows from individual evaluations into purchase support and institutional software.'
],[('Primary sector:','Automotive technology'),('Secondary sector or Sub-sector:','Consumer buying software; business software and licensing')])
row(T,'Goals',[
'At the Hurricane Pitch Competition, Curbside seeks early funding, pilot customers, advisors, and introductions to inspection providers, credit unions, and automotive platforms. These relationships would help validate both the consumer product and the broader buying-platform opportunity.',
'Proposed milestones include 20 buyer interviews and a 25-person pilot within 30 days; initial paid purchases and ten prospective partner interviews within 90 days; and a defined service-provider or institutional pilot within six months. The first-year consumer target is 500 paying customers.',
'An illustrative $10,000 award would support $3,000 of evaluation reliability work, $2,500 of customer research and acquisition, $2,500 of pilot development, and $2,000 of specialist review and contingency. The funding request and allocations require founder confirmation.'
])
p=D.add_paragraph('Financial assumptions: 2026–2028 are illustrative consumer projections, using 500, 2,000, and 5,000 annual payers at $12 average revenue per payer. Expenditures exclude founder salary, taxes, and incremental institutional development. Partnership and licensing revenue remain unmodeled until pilots establish terms. Historical figures require confirmation.');p.paragraph_format.space_before=Pt(7)
for r in p.runs:r.font.size=Pt(9)
t=D.add_table(rows=1,cols=6);t.style='Table Grid'
for c,v in zip(t.rows[0].cells,['Financials ($ US)','2024','2025','2026 (projected)','2027 (projected)','2028 (projected)']):c.text=v
for vals in [('Revenues','[Confirm]','[Confirm]','6,000','24,000','60,000'),('Expenditures','[Confirm]','[Confirm]','12,000','24,000','42,000'),('Net','[Confirm]','[Confirm]','(6,000)','0','18,000')]:
 for c,v in zip(t.add_row().cells,vals):c.text=v
for rowx in t.rows:
 for c in rowx.cells:
  for p in c.paragraphs:
   p.paragraph_format.space_after=Pt(2)
   for r in p.runs:r.font.size=Pt(8)
p=D.add_paragraph('[1] Internal calibration review, docs/calibration-pass-1.md. [2] Cox Automotive, December 2025, 2026 Forecasts. [3] kbb.com/whats-my-car-worth/ [4] cargurus.com/research/car-valuation [5] carfax.com/vehicle-history-reports/ [6] Curbside pricing configuration. Sources reviewed September 29, 2026.')
for r in p.runs:r.font.size=Pt(8)
D.save('output/documents/Curbside_Executive_Summary_Template.docx')
