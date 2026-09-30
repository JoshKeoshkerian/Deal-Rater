from docx import Document
D=Document('tmp/executive-summary/before-trim.docx');p=D.paragraphs
edits={
4:'Founded in 2026, Curbside is building a platform to guide private-party used-car buyers from comparing vehicles and negotiating to arranging inspections, financing, and secure transactions.',
5:'Our first product is a Chrome extension that compares asking prices directly across Facebook Marketplace listings. It combines comparable prices, vehicle risk indicators, and negotiation guidance in an explained score out of 100.',
6:'The broader vision is a connected buying experience supported by service-provider partnerships, with software subscriptions and licensing for credit unions and automotive platforms extending our reach and creating recurring business revenue.',
8:'Private-party buyers make a major financial decision with incomplete information. Listings can omit important details, asking prices are difficult to compare, and inexperienced buyers may struggle to recognize warning signs or negotiate.',
9:'Buyers must also coordinate vehicle history, inspections, financing, and purchase logistics across separate services. This burden particularly affects students, first-time buyers, and budget-conscious households that depend on reliable transportation.',
10:'Curbside offers faster screening and practical next steps. Free evaluations introduce the product; paid purchases and customer interviews will test willingness to pay.',
13:'The extension compares a vehicle with other Facebook Marketplace asking prices, adjusting for characteristics such as mileage and trim. Saved evaluations help buyers build a shortlist. Asking prices are not completed sale prices, and evaluations do not replace inspections.',
14:'Planned services would let buyers prepare seller questions, arrange inspections, and connect with financing and secure transaction providers. Institutional integrations could place these tools within a credit union’s member experience or an automotive platform. Provider relationships are not yet confirmed.',
16:'Implemented capabilities include listing capture, evaluation, accounts, saved results, and payment integration code. Internal testing found 15.3% median absolute percentage error across 4,133 held-out asking-price predictions and 38% overall verdict agreement across 72 labels from one reviewer. Independent validation remains necessary. [1]',
17:'Curbside owns its intellectual property, including its original software, evaluation methods, and proprietary know-how.',
18:'Operating priorities include platform permissions, privacy, data rights, and specialist review of future financing and transaction arrangements. Referral compensation must be disclosed and kept separate from evaluation scores.',
20:'The initial market is U.S. private-party buyers shopping on Facebook Marketplace through desktop Chrome. Cox Automotive projected 38.3 million U.S. used-vehicle sales for 2026, a broad industry benchmark rather than a count of eligible users. [2]',
21:'An illustrative reachable segment of 5% would represent 1.915 million shopping episodes; this assumption needs validation. Institutional demand and contract values will be assessed through pilot discussions.',
23:'Priority consumers are students, first-time buyers, and budget-conscious households seeking reliable transportation. They need clear comparisons and help moving from a listing to a purchase.',
24:'Future business customers include credit unions and automotive platforms. Inspection, financing, and transaction providers are potential service partners.',
26:'Distribution will begin through the Chrome Web Store and Curbside website. Demonstrations will show direct comparisons with other Facebook Marketplace listings. Educational content, campus outreach, automotive communities, and referrals will drive discovery; paid acquisition will expand only when conversion and acquisition costs justify it.',
27:'Conservative assumes weak organic reach and conversion; realistic assumes steady content, campus outreach, and referrals; optimistic assumes stronger referrals and repeatable community distribution. All three are planning assumptions.',
28:'Qualified visits are prospective-buyer website or store visits; activation is a completed first evaluation. Paying customers = activated shoppers × paid conversion. Counts assume unique shoppers within each year. Values below follow 2026 / 2027 / 2028; 2026 is the launch-year total.',
29:'Budgets cover content, outreach, and paid tests; founder time is unpaid. Traffic depends mainly on organic distribution. Conversion and acquisition costs will be reviewed monthly. No signed distribution agreements or partner-generated traffic are assumed.',
32:'Consumer freemium pricing includes ten free checks, ten additional checks for $7.99, twenty for $11.99, or unlimited evaluations for $16.99 per month. Paid packs do not expire. [6]',
33:'Planned revenue streams include referral partnerships for inspection, financing, and transaction services, plus recurring software subscriptions and licensing for credit unions and automotive platforms. Terms will be established through pilots.',
34:'Curbside has no revenue or funding to date and has incurred no significant operating costs. The projections below model consumer adoption and spending; they exclude founder compensation and institutional development.',
36:'Kelley Blue Book offers private-party valuations, CarGurus provides valuation tools and deal ratings, and CARFAX supplies vehicle-history reports. [3–5] CarScout directly competes through Marketplace comparable-listing searches and ranked alternatives. [7]',
37:'Free research tools and manual comparisons create a price hurdle. Curbside must demonstrate enough convenience and useful guidance to earn paid purchases. Pilots will compare usability, trust, and willingness to pay against existing tools.',
42:'Curbside combines direct Facebook Marketplace price comparisons with risk explanations and negotiation guidance. Its broader opportunity is to carry that vehicle context through inspection, financing, and transaction steps, reducing the research buyers must repeat across services.',
43:'Long-term differentiation would come from permitted buyer and inspection feedback, service integrations, and distribution through trusted institutions. These advantages remain to be developed; the current feature set is reproducible, with no exclusive network or proven data advantage.',
48:'The founders combine business analysis, software development, and engineering. University networks support early buyer research; automotive advisors and specialists in partnerships, security, and integrations will help address gaps as the platform expands.',
50:'Curbside seeks $7,500 in early funding, pilot customers, advisors, and introductions to inspection providers, credit unions, and automotive platforms.',
51:'Milestones are 20 buyer interviews and a 25-person pilot within 30 days; initial paid purchases and ten partner interviews within 90 days; and a defined provider or institutional pilot within six months. Launch-year payer targets are 60 conservative, 500 realistic, and 1,600 optimistic, subject to launch timing.',
52:'The proposed $7,500 allocation is $2,500 for evaluation reliability and product improvements, $2,000 for customer research and acquisition, $1,500 for partner-pilot development, and $1,500 for specialist advice and contingency. These fund validation rather than additional recurring annual expenses.'
}
for i,t in edits.items():
 # Keep existing paragraph font sizes.
 sizes=[r.font.size for r in p[i].runs];p[i].text=t
 if sizes and sizes[0]:p[i].runs[0].font.size=sizes[0]
for i in [11,15,30,38,39,40,44]:p[i]._element.getparent().remove(p[i]._element)
for x in D.paragraphs:
 if x.text!='Financials ($ US)':x.paragraph_format.page_break_before=False
D.save('output/documents/Curbside_Executive_Summary.docx')
