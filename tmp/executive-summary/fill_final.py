from docx import Document
p='output/documents/Curbside_Executive_Summary.docx';d=Document(p)
replacements={
'[Contact name and cell phone]':'Joshua Keoshkerian | (314) 520-9265',
'Our intellectual property position centers on original software, evaluation methods, and proprietary know-how. [Confirm copyright ownership, trademark or patent status, and third-party licensing.]':'Curbside owns its intellectual property, including its original software, evaluation methods, and proprietary know-how.',
'Current revenue, monthly burn, and capital received: [confirm actual amounts and sources, or none].':'Curbside has not yet generated revenue or received funding and has incurred no significant operating costs to date.',
'[confirm funding amount]':'$7,500',
'An illustrative $10,000 award would support $3,000 of evaluation reliability work, $2,500 of customer research and acquisition, $2,500 of pilot development, and $2,000 of specialist review and contingency. The funding request and allocations require founder confirmation.':'The proposed $7,500 allocation is $2,500 for evaluation reliability and product improvements, $2,000 for customer research and acquisition, $1,500 for initial partner-pilot development, and $1,500 for specialist advice and contingency. These uses support the next stage of validation; they are not additional recurring annual expenses.'
}
for x in d.paragraphs:
 for a,b in replacements.items():
  if a in x.text:
   # Preserve run formatting where possible.
   if any(a in r.text for r in x.runs):
    for r in x.runs:
     if a in r.text:r.text=r.text.replace(a,b)
   else:x.text=x.text.replace(a,b)
d.save(p)
import re
assert not any(re.search(r'\[(?:confirm|Contact|Founder|Identify)',x.text,re.I) for x in d.paragraphs)
