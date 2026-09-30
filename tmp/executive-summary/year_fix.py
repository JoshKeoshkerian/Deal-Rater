from docx import Document
p='output/documents/Curbside_Executive_Summary.docx'
d=Document(p)
for x in d.paragraphs:
 if x.text.startswith('Curbside is building'):
  x.text=x.text.replace('Curbside is building','Started in 2026, Curbside is building',1)
 if x.text.startswith('Illustrative consumer projections'):
  for r in x.runs:
   r.text=r.text.replace('Illustrative consumer projections assume','Curbside started in 2026. Illustrative projections for launch year 2026 and calendar years 2027 and 2028 assume').replace('Historical figures require confirmation.','These are draft scenarios, not actual results; the launch-year assumptions require confirmation against the 2026 operating period.')
t=d.tables[0]
for row in t.rows:
 for idx in [2,1]:
  c=row.cells[idx];row._tr.remove(c._tc)
for idx in [2,1]:t._tbl.tblGrid.remove(t._tbl.tblGrid.gridCol_lst[idx])
d.save(p)
